import requests
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
import time
import sys

ACCESS_TOKEN = "EAAc41mHwnzcBRrljeJEOga1jQ1vE4IBtZBfUSpY2UFjs0LTRuYfJwVQCSZBIkvRYjqHTyxZC0CsNaWVZCZAgkiAJvGgerZBaNAZCeEa266hzehXCxtwFN2UujjgUgWrCsOgXKLxnjNd1sEy8dKaabnoZCQnmbTNQBV5WrePGuHt10C3TgSsYaO2rvUW7vNmpfMQZD"
BASE_URL = "https://graph.facebook.com/v19.0/ads_archive"

KEYWORDS = [
    "rhinoplasty", "facelift", "breast augmentation", "breast implants",
    "tummy tuck", "liposuction", "mommy makeover", "blepharoplasty",
    "body contouring", "BBL", "neck lift", "chin implant", "otoplasty",
    "plastic surgeon", "plastic surgery", "facial plastic surgery",
    "cosmetic surgeon", "cosmetic surgery", "aesthetic surgeon"
]

LA_TERMS = [
    "los angeles", "beverly hills", "santa monica", "west hollywood",
    "brentwood", "pacific palisades", "calabasas", "encino",
    "sherman oaks", "studio city", "burbank", "pasadena",
    "long beach", "manhattan beach", "culver city", "venice",
    "marina del rey", "southern california", "socal", "90210"
]

EXCLUDE_TERMS = [
    "skin care brand", "serum", "supplement", "medical device",
    "equipment", "software", "school", "training course",
    "insurance", "attorney", "hair salon"
]

FIELDS = (
    "page_name,page_id,ad_creative_bodies,ad_creative_link_captions,"
    "ad_creative_link_titles,ad_delivery_start_time,ad_snapshot_url,publisher_platforms"
)

MAX_PAGES = 5


def fetch_ads_for_keyword(keyword):
    ads = []
    params = {
        "access_token": ACCESS_TOKEN,
        "ad_reached_countries": "US",
        "ad_active_status": "ACTIVE",
        "search_terms": keyword,
        "limit": 100,
        "fields": FIELDS,
    }
    url = BASE_URL
    page_count = 0

    while url and page_count < MAX_PAGES:
        try:
            resp = requests.get(url, params=params if page_count == 0 else {}, timeout=30)
            data = resp.json()
        except Exception as e:
            print(f"  Network error for '{keyword}': {e}")
            break

        if "error" in data:
            err = data["error"]
            code = err.get("code")
            msg = err.get("message", "")
            if code in (190, 102, 104) or "token" in msg.lower() or "permission" in msg.lower():
                print(f"\n=== TOKEN / PERMISSION ERROR ===")
                print(f"Code: {code}")
                print(f"Message: {msg}")
                print("\nQué hacer:")
                print("1. Ve a https://developers.facebook.com/tools/explorer/")
                print("2. Genera un nuevo User Access Token con permisos: ads_read")
                print("3. Reemplaza ACCESS_TOKEN en este script con el nuevo token.")
                sys.exit(1)
            else:
                print(f"  API error for '{keyword}': {err}")
                break

        results = data.get("data", [])
        ads.extend(results)
        page_count += 1

        next_url = data.get("paging", {}).get("next")
        if next_url:
            url = next_url
            params = {}
        else:
            break

    return ads


def normalize_text(text):
    if isinstance(text, list):
        return " ".join(str(t) for t in text).lower()
    return str(text).lower() if text else ""


def contains_any(text, terms):
    t = normalize_text(text)
    return any(term in t for term in terms)


def main():
    print("=" * 60)
    print("Meta Ads Library — LA Plastic Surgeons Extractor")
    print("=" * 60)

    all_ads = []
    for kw in KEYWORDS:
        print(f"Buscando: '{kw}' ...", end=" ", flush=True)
        ads = fetch_ads_for_keyword(kw)
        print(f"{len(ads)} anuncios")
        all_ads.extend(ads)
        time.sleep(0.5)

    print(f"\nTotal raw (con duplicados):       {len(all_ads)}")

    # Deduplicate by page_id — keep first occurrence
    seen = {}
    for ad in all_ads:
        pid = ad.get("page_id")
        if pid and pid not in seen:
            seen[pid] = ad

    deduped = list(seen.values())
    print(f"Después de deduplicar por page_id: {len(deduped)}")

    # Filter: must mention LA
    la_filtered = []
    for ad in deduped:
        page_name = ad.get("page_name", "")
        bodies = ad.get("ad_creative_bodies", [])
        captions = ad.get("ad_creative_link_captions", [])
        titles = ad.get("ad_creative_link_titles", [])
        combined = " ".join([
            normalize_text(page_name),
            normalize_text(bodies),
            normalize_text(captions),
            normalize_text(titles),
        ])
        if contains_any(combined, LA_TERMS):
            la_filtered.append(ad)

    print(f"Después de filtro LA:              {len(la_filtered)}")

    # Exclude non-surgeons
    final = []
    for ad in la_filtered:
        page_name = ad.get("page_name", "")
        bodies = ad.get("ad_creative_bodies", [])
        combined = normalize_text(page_name) + " " + normalize_text(bodies)
        if not contains_any(combined, EXCLUDE_TERMS):
            final.append(ad)

    print(f"Después de filtro no-cirujanos:    {len(final)}")

    # Build DataFrame
    rows = []
    for ad in final:
        page_name = ad.get("page_name", "")
        page_id = ad.get("page_id", "")
        bodies = ad.get("ad_creative_bodies", [])
        ad_text = " ".join(str(b) for b in bodies) if isinstance(bodies, list) else str(bodies)
        ad_preview = ad_text[:300] if ad_text else ""
        platforms = ad.get("publisher_platforms", [])
        platforms_str = ", ".join(platforms) if isinstance(platforms, list) else str(platforms)
        rows.append({
            "Page Name": page_name,
            "Page ID": page_id,
            "Facebook Page URL": f"https://www.facebook.com/{page_id}",
            "Ad Text Preview": ad_preview,
            "Platforms": platforms_str,
            "Ad Start Date": ad.get("ad_delivery_start_time", ""),
            "Ad Snapshot URL": ad.get("ad_snapshot_url", ""),
            "Instagram Handle": "",
            "Website": "",
            "Email": "",
            "Notes": "",
        })

    df = pd.DataFrame(rows)
    df.sort_values("Page Name", inplace=True)
    df.reset_index(drop=True, inplace=True)

    output_file = "la_plastic_surgeons_ads.xlsx"
    df.to_excel(output_file, index=False)

    # Format Excel
    wb = load_workbook(output_file)
    ws = wb.active

    header_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="276221")
    cell_font = Font(name="Arial", size=10)

    col_widths = {
        "A": 40,  # Page Name
        "B": 18,  # Page ID
        "C": 42,  # Facebook Page URL
        "D": 55,  # Ad Text Preview
        "E": 22,  # Platforms
        "F": 15,  # Ad Start Date
        "G": 45,  # Ad Snapshot URL
        "H": 22,  # Instagram Handle
        "I": 28,  # Website
        "J": 28,  # Email
        "K": 30,  # Notes
    }

    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = cell_font
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    ws.freeze_panes = "A2"

    wb.save(output_file)
    print(f"\nArchivo generado: {output_file}  ({len(df)} anunciantes)")
    print("=" * 60)


if __name__ == "__main__":
    main()
